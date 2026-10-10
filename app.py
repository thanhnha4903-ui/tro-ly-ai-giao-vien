import streamlit as st
import google.generativeai as genai
import re
import urllib.parse
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# =====================================================================
# 1. THIẾT LẬP TRANG & CSS
# =====================================================================
st.set_page_config(page_title="Trợ Lý AI Giáo Viên Pro", layout="wide", page_icon="👨‍🏫")

if "ket_qua_ai" not in st.session_state:
    st.session_state.ket_qua_ai = ""
if "danh_sach_anh" not in st.session_state:
    st.session_state.danh_sach_anh = {}

st.markdown("""
<style>
    h1 { text-align: center; color: #0284c7 !important; text-shadow: 0px 2px 4px rgba(0,0,0,0.1); font-family: 'Segoe UI', Tahoma, sans-serif;}
    .stTabs [data-baseweb="tab-list"] { gap: 10px; }
    .stTabs [data-baseweb="tab"] { border-radius: 5px 5px 0 0; background-color: #f0f9ff; border: 1px solid #bae6fd; }
    .stButton>button[kind="primary"] { background: linear-gradient(90deg, #0284c7 0%, #0ea5e9 100%); color: white; border-radius: 8px; width: 100%; font-weight: bold; font-size: 16px; padding: 10px; border: none; transition: 0.3s; }
    .stButton>button[kind="primary"]:hover { transform: translateY(-3px); box-shadow: 0 5px 15px rgba(14, 165, 233, 0.4); }
    .a4-paper { background-color: white; padding: 40px; border-radius: 8px; box-shadow: 0 4px 6px rgba(0,0,0,0.05); color: #333; border: 1px solid #e2e8f0; }
    .api-guide-box { background-color: #f8fafc; border-left: 4px solid #f59e0b; padding: 15px; border-radius: 4px; font-size: 14px; color: #334155; margin-bottom: 20px;}
    .api-guide-box a { color: #0284c7; font-weight: bold; text-decoration: none;}
    .api-guide-box a:hover { text-decoration: underline;}
</style>
""", unsafe_allow_html=True)

# =====================================================================
# HÀM HỖ TRỢ: TẠO FILE WORD ĐẸP, CHUẨN SƯ PHẠM
# =====================================================================
def tao_file_word_chuan(noi_dung, ten_file, loai_tai_lieu, mon_hoc, lop, ten_bai):
    doc = Document()
    
    # Thiết lập Font chữ mặc định (Times New Roman, Size 13 - Chuẩn văn bản hành chính)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)

    # 1. TẠO TIÊU ĐỀ
    if "Kế hoạch bài dạy" in loai_tai_lieu or "Đề kiểm tra" in loai_tai_lieu:
        p_header = doc.add_paragraph()
        p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_header = p_header.add_run(f"{'KẾ HOẠCH BÀI DẠY (GIÁO ÁN)' if 'Kế hoạch' in loai_tai_lieu else 'ĐỀ KIỂM TRA'}\nMÔN: {mon_hoc.upper()} {lop}\n")
        run_header.bold = True
        run_header.font.size = Pt(14)
        
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(f"BÀI: {ten_bai.upper()}")
        run_title.bold = True
        run_title.font.size = Pt(16)
    else:
        doc.add_heading(ten_file, level=0)

    # 2. XỬ LÝ NỘI DUNG TỪ MARKDOWN SANG WORD ĐẸP
    lines = noi_dung.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Xử lý các Tiêu đề chính (Bắt đầu bằng I, II, III, IV...)
        if re.match(r'^[IIVX]+\.', line) or line.startswith('#'):
            clean_line = line.replace('#', '').strip()
            clean_line = clean_line.replace('**', '') 
            p = doc.add_paragraph()
            run = p.add_run(clean_line)
            run.bold = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            
        # Xử lý các mục nhỏ (1., 2., 3. hoặc a., b., c.)
        elif re.match(r'^\d+\.', line) or re.match(r'^[a-z]\.', line):
            clean_line = line.replace('**', '').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(clean_line)
            run.bold = True
            p.paragraph_format.space_after = Pt(4)
            
        # Xử lý gạch đầu dòng (Chuyển thành dấu - chuẩn, lùi đầu dòng, bỏ dấu *)
        elif line.startswith('-') or line.startswith('*') or line.startswith('+'):
            clean_line = line[1:].strip().replace('**', '')
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.add_run(f"– {clean_line}") 
            p.paragraph_format.space_after = Pt(4)
            
        # Xử lý văn bản bình thường
        else:
            clean_line = line.replace('**', '') 
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.5) # Thụt lề dòng đầu tiên
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY # Căn đều 2 bên
            p.add_run(clean_line)
            p.paragraph_format.space_after = Pt(6)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# =====================================================================
# 2. THANH BÊN (SIDEBAR)
# =====================================================================
with st.sidebar:
    st.image("https://cdn-icons-png.flaticon.com/512/1903/1903162.png", width=80)
    st.header("⚙️ Bảng Điều Khiển")
    
    st.markdown("""
    <div class="api-guide-box">
        <strong>💡 Hướng dẫn lấy API Key (Miễn phí):</strong><br>
        1. Truy cập trang <a href="https://aistudio.google.com/app/apikey" target="_blank">Google AI Studio</a>.<br>
        2. Đăng nhập bằng tài khoản Google.<br>
        3. Nhấn nút <strong>Create API Key</strong> màu xanh.<br>
        4. Copy dãy ký tự vừa tạo và dán vào ô bên dưới.
    </div>
    """, unsafe_allow_html=True)
    
    api_key_input = st.text_input("🔑 Nhập API Key của Thầy/Cô:", type="password")
    
    st.markdown("---")
    st.markdown("""
    **✅ Hỗ trợ Toán học:** Tương thích MathType.
    **✅ Hình ảnh tự động:** Lên kịch bản tạo ảnh.
    **✅ Xuất File Word:** Trình bày đẹp chuẩn sư phạm.
    """)

if api_key_input:
    genai.configure(api_key=api_key_input)
    model = genai.GenerativeModel('gemini-3.8-flash')

# =====================================================================
# 3. GIAO DIỆN CHÍNH
# =====================================================================
st.title("👨‍🏫 TRỢ LÝ AI ĐA NĂNG CHO GIÁO VIÊN")
st.markdown("<p style='text-align: center; color: #64748b; font-size: 18px;'>Soạn Giáo Án - Lên Kịch Bản Slide - Tạo Đề Kiểm Tra trong tích tắc</p>", unsafe_allow_html=True)
st.markdown("---")

loai_tai_lieu = st.radio(
    "🎯 Thầy/Cô muốn tạo tài liệu gì hôm nay?",
    ["📝 Kế hoạch bài dạy (CV 5512)", "📽️ Kịch bản Slide bài giảng", "📋 Đề kiểm tra & Đáp án"],
    horizontal=True
)
st.markdown("<br>", unsafe_allow_html=True)

col1, col2 = st.columns(2)
with col1:
    mon_hoc = st.selectbox("📚 Môn học", ["Toán học", "Ngữ Văn", "Tiếng Anh", "Vật Lý", "Hóa học", "Sinh học", "Lịch sử", "Địa lý", "Tin học", "Công nghệ", "Khác"])
    lop = st.selectbox("🎓 Khối Lớp", ["Lớp 6", "Lớp 7", "Lớp 8", "Lớp 9", "Lớp 10", "Lớp 11", "Lớp 12"])
    
with col2:
    ten_bai = st.text_input("📝 Tên bài học / Chủ đề kiểm tra", placeholder="VD: Định lý Vi-ét, Sinh sản ở thực vật...")
    thoi_luong = st.number_input("⏱️ Thời lượng (Số tiết / Số phút)", min_value=1, value=1)

with st.expander("🛠️ Tùy chỉnh chi tiết (Không bắt buộc)"):
    if loai_tai_lieu == "📋 Đề kiểm tra & Đáp án":
        cau_truc_de = st.text_input("Cấu trúc đề (VD: 20 câu trắc nghiệm, 2 câu tự luận):")
    else:
        phuong_phap = st.selectbox("Phương pháp dạy học", ["Linh hoạt", "Dạy học dự án", "Lớp học đảo ngược", "Thảo luận nhóm", "Trực quan sinh động"])
        trinh_do = st.selectbox("Trình độ học sinh", ["Đại trà (Khá - Trung bình)", "Học sinh Giỏi / Chuyên", "Cần kèm cặp thêm"])
    yeu_cau_them = st.text_area("Ghi chú đặc biệt cho AI:")

# =====================================================================
# 4. XỬ LÝ LỆNH (PROMPT)
# =====================================================================
if st.button("🚀 TIẾN HÀNH TẠO TÀI LIỆU", type="primary"):
    if not api_key_input:
        st.error("⚠️ Thầy/Cô vui lòng lấy mã API Key theo hướng dẫn bên trái và dán vào ô nhé!")
    elif not ten_bai:
        st.warning("⚠️ Thầy/Cô vui lòng nhập Tên bài học / Chủ đề kiểm tra.")
    else:
        with st.spinner("AI đang soạn thảo tài liệu chất lượng cao... Thầy/Cô đợi một chút nhé..."):
            
            lenh_hinh_anh = """
            Nếu cần hình ảnh minh họa, BẮT BUỘC dùng cú pháp sau (AI không vẽ hình trực tiếp, mà ghi chú để tạo ảnh sau):
            [Ảnh minh họa X: <câu lệnh tạo ảnh bằng tiếng Anh để ảnh chính xác nhất>]
            """
            
            prompt_han_che_bullet = """
            LƯU Ý TRÌNH BÀY QUAN TRỌNG:
            - Trình bày như một giáo viên lâu năm: Hãy diễn đạt thành các đoạn văn hoàn chỉnh, mạch lạc.
            - HẠN CHẾ TỐI ĐA việc lạm dụng gạch đầu dòng (*, -, +). CHỈ dùng gạch đầu dòng khi thực sự cần liệt kê ngắn gọn.
            - KHÔNG sử dụng ký hiệu in đậm Markdown (**) thừa thãi.
            """

            if loai_tai_lieu == "📝 Kế hoạch bài dạy (CV 5512)":
                prompt = f"""
                Bạn là chuyên gia sư phạm. Lập Kế hoạch bài dạy theo chuẩn Công văn 5512.
                Môn: {mon_hoc} {lop}. Bài: "{ten_bai}". Trình độ: {trinh_do}. Phương pháp: {phuong_phap}.
                - Gồm: I. Mục tiêu, II. Thiết bị, III. Tiến trình dạy học (4 hoạt động).
                - Công thức Toán/Lý/Hóa bọc trong $$ hoặc $ (chuẩn LaTeX).
                {lenh_hinh_anh}
                {prompt_han_che_bullet}
                """
            elif loai_tai_lieu == "📽️ Kịch bản Slide bài giảng":
                prompt = f"""
                Bạn là chuyên gia sư phạm. Viết kịch bản chi tiết thiết kế Slide PowerPoint.
                Môn: {mon_hoc} {lop}. Bài: "{ten_bai}". Thời lượng: {thoi_luong} tiết.
                - Mỗi Slide nêu rõ: Nội dung text hiển thị, Ghi chú cho giáo viên thuyết trình.
                - Công thức Toán bọc trong $$ hoặc $.
                {lenh_hinh_anh}
                {prompt_han_che_bullet}
                """
            else:
                prompt = f"""
                Bạn là chuyên gia ra đề thi. Tạo Đề kiểm tra và Đáp án chi tiết.
                Môn: {mon_hoc} {lop}. Chủ đề: "{ten_bai}". Thời gian: {thoi_luong} phút. Cấu trúc: {cau_truc_de if 'cau_truc_de' in locals() else 'Tự luận'}.
                - Gồm: I. Ma trận, II. Đề thi, III. Đáp án.
                - Công thức Toán bọc trong $$.
                {lenh_hinh_anh}
                {prompt_han_che_bullet}
                """

            try:
                response = model.generate_content(prompt)
                st.session_state.ket_qua_ai = response.text
                
                pattern = r"\[(Ảnh minh họa \d+):\s*(.*?)\]"
                matches = re.findall(pattern, response.text)
                st.session_state.danh_sach_anh = {match[0]: match[1] for match in matches}
                
                st.rerun() 
            except Exception as e:
                st.error(f"Có lỗi từ máy chủ AI: {e}. Vui lòng thử lại.")

# =====================================================================
# 5. HIỂN THỊ KẾT QUẢ & XUẤT WORD
# =====================================================================
if st.session_state.ket_qua_ai:
    st.markdown("---")
    tab1, tab2 = st.tabs(["📄 Nội dung Tài liệu", "🎨 Trạm tạo Ảnh minh họa"])
    
    with tab1:
        st.success("🎉 Tạo thành công! Thầy/Cô hãy bấm nút Tải xuống File Word (.docx) bên dưới nhé.")
        
        ten_file_xuat = f"{loai_tai_lieu.split(' ')[1]}_{mon_hoc}_{ten_bai}"
        
        file_word_bytes = tao_file_word_chuan(st.session_state.ket_qua_ai, ten_file_xuat, loai_tai_lieu, mon_hoc, lop, ten_bai)
        
        st.download_button(
            label="📄 TẢI XUỐNG FILE WORD (.docx) CHUẨN SƯ PHẠM", 
            data=file_word_bytes, 
            file_name=f"{ten_file_xuat}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
            type="secondary"
        )
        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown(f"<div class='a4-paper'>{st.session_state.ket_qua_ai}</div>", unsafe_allow_html=True)
        
    with tab2:
        st.subheader("🎨 Hệ thống Tự động vẽ ảnh")
        st.info("Dựa vào nội dung bài giảng, AI đã ghi chú các vị trí cần hình ảnh và viết sẵn câu lệnh tạo ảnh.")
        
        if st.session_state.danh_sach_anh:
            anh_chon = st.selectbox("📌 Chọn vị trí ảnh muốn tạo:", list(st.session_state.danh_sach_anh.keys()))
            prompt_ve_anh = st.session_state.danh_sach_anh[anh_chon]
            prompt_da_sua = st.text_area("📝 Câu lệnh tiếng Anh tạo ảnh (Thầy/Cô có thể sửa đổi):", value=prompt_ve_anh, height=80)
            
            if st.button(f"🖌️ Tiến hành vẽ [{anh_chon}]", type="primary"):
                with st.spinner("Đang vẽ ảnh sắc nét..."):
                    encoded_prompt = urllib.parse.quote(prompt_da_sua)
                    image_url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=800&height=600&nologo=true"
                    st.image(image_url, caption=f"Đã tạo xong: {anh_chon}", use_container_width=True)
                    st.success("💡 Mẹo: Thầy/Cô hãy Click chuột phải vào ảnh -> Chọn 'Sao chép hình ảnh' và dán thẳng vào File Word.")
        else:
            st.warning("Bài học này AI không cần hình ảnh phức tạp.")
